#!/usr/bin/env python3
"""
Importa a planilha "Sabor futebol.xlsx" (aba "Rachão") para server/data/db.json.

Uso:
    python3 tools/importar_planilha.py "Sabor futebol.xlsx" [server/data/db.json]

Requer: pip install openpyxl

Layout esperado da aba (o mesmo da planilha original):
  - Linha 1: "Valor do jogo:" + valor, a cada par de colunas (Jogou? / Pagou?)
  - Linha 2: data de cada jogo (coluna "Jogou?")
  - Linhas 5+: um atleta por linha, coluna A = nome, "X" marca jogou/pagou
  - A1/A2: "Chave pix:" e a chave
Os totais da planilha (colunas B, C, D) NÃO são importados: o app recalcula tudo
a partir de quem jogou e quem pagou em cada jogo.
"""
import json
import re
import secrets
import sys
import uuid
from pathlib import Path

import openpyxl

ALFABETO = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # sem 0/O, 1/I para evitar confusão


def novo_codigo(usados):
    while True:
        c = "".join(secrets.choice(ALFABETO) for _ in range(6))
        if c not in usados:
            usados.add(c)
            return c


def comentario(cell):
    if not cell.comment:
        return ""
    txt = cell.comment.text
    m = re.search(r"Comment:\s*(.*)", txt, re.S)
    return (m.group(1) if m else txt).strip()


def main():
    origem = Path(sys.argv[1] if len(sys.argv) > 1 else "Sabor futebol.xlsx")
    destino = Path(sys.argv[2] if len(sys.argv) > 2 else "server/data/db.json")

    wb = openpyxl.load_workbook(origem, data_only=True)
    ws = wb["Rachão"] if "Rachão" in wb.sheetnames else wb.worksheets[0]

    # --- jogos (colunas E, G, I, ... = "Jogou?"; a seguinte é "Pagou?")
    colunas = []
    col = 5
    while col <= ws.max_column:
        data = ws.cell(2, col).value
        valor = ws.cell(1, col + 1).value
        if data is None:
            break
        colunas.append((col, data.date().isoformat(), float(valor)))
        col += 2

    # --- atletas (linha 5 até antes de "Totais")
    linhas = []
    r = 5
    while r <= ws.max_row:
        nome = ws.cell(r, 1).value
        if nome is None or str(nome).strip().lower() == "totais":
            break
        linhas.append(r)
        r += 1

    usados = set()
    atletas = []
    id_por_linha = {}
    for r in linhas:
        a = {
            "id": str(uuid.uuid4()),
            "nome": str(ws.cell(r, 1).value).strip(),
            "codigo": novo_codigo(usados),
            "observacoes": comentario(ws.cell(r, 1)),
        }
        atletas.append(a)
        id_por_linha[r] = a["id"]

    eventos = []
    for col, data, valor in colunas:
        participantes = [id_por_linha[r] for r in linhas if ws.cell(r, col).value == "X"]
        pagamentos = [id_por_linha[r] for r in linhas if ws.cell(r, col + 1).value == "X"]
        assert set(pagamentos) <= set(participantes), f"Pagamento sem participação em {data}"
        eventos.append(
            {
                "id": str(uuid.uuid4()),
                "data": data,
                "valorTotal": valor,
                "participantes": participantes,
                "pagamentos": pagamentos,
            }
        )
    eventos.sort(key=lambda e: e["data"])

    pix = ws["A2"].value or ""
    db = {
        "grupo": {"nome": "saBORR futebol", "chavePix": str(pix).strip()},
        "atletas": atletas,
        "eventos": eventos,
    }
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(json.dumps(db, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{len(atletas)} atletas e {len(eventos)} jogos gravados em {destino}")


if __name__ == "__main__":
    main()
